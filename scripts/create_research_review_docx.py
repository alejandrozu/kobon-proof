"""Create a restrained, readable Word edition of the cumulative review."""
from pathlib import Path
import re,copy,os,subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
(ROOT/'work').mkdir(exist_ok=True)
source=ROOT/'research/six-hour-2026-09-21/RESEARCH_REVIEW.md'
md=source.read_text(encoding='utf-8')
# LibreOffice's OMML floor glyph fallback is ambiguous on this host. Standard
# named floor/ceil operators preserve editable mathematics and render clearly.
math_input=re.sub(r'(?:\\left\s*)?\\lfloor(.*?)(?:\\right\s*)?\\rfloor',
    r'\\operatorname{floor}\\left(\1\\right)',md,flags=re.S)
math_input=re.sub(r'(?:\\left\s*)?\\lceil(.*?)(?:\\right\s*)?\\rceil',
    r'\\operatorname{ceil}\\left(\1\\right)',math_input,flags=re.S)
subprocess.run([os.environ.get('PANDOC_EXE','pandoc'),'--from','markdown+tex_math_single_backslash',
    '-o',str(ROOT/'work/review-math.docx')],input=math_input,text=True,encoding='utf-8',check=True)
math_doc=Document(ROOT/'work/review-math.docx')
equations=math_doc._element.xpath('//m:oMathPara')
doc=Document();sec=doc.sections[0]
doc.settings.odd_and_even_pages_header_footer=False
sec.different_first_page_header_footer=False
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=sec.footer_distance=Inches(.3)
for sty in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
    s=doc.styles[sty];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].paragraph_format.space_after=Pt(12)
for st,size in [('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    doc.styles[st].font.size=Pt(size)
    doc.styles[st].paragraph_format.space_before=Pt(13)
    doc.styles[st].paragraph_format.space_after=Pt(6)
head=sec.header.paragraphs[0]
head.text='Kobon triangle research  |  Alejandro Zarzuelo Urdiales'
head.runs[0].font.size=Pt(9)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Research review  •  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for r in foot.runs:r.font.size=Pt(9)

def link(p,label,url):
    if not url.startswith('http'):
        url='https://github.com/alejandrozu/kobon-proof/blob/main/'+url
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1F4E79');pr.append(c);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,s):
    for token in re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)',s):
        if not token:continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if m:link(p,m[1],m[2]);continue
        if token.startswith('**'):
            p.add_run(token[2:-2]).bold=True
        elif token.startswith('`'):
            r=p.add_run(token[1:-1]);r.font.name='Consolas';r.font.size=Pt(9.5)
        else:p.add_run(re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)',r'\1',token))

def table(rows):
    rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)]
    data=[[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    t=doc.add_table(rows=0,cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    cols=len(data[0]);widths=([.75,2.0,2.0] if cols==3 and data[0][0]=='Lines' else ([2.0,4.9] if cols==2 else [6.9/cols]*cols))
    if len(widths)!=cols:widths=[6.9/cols]*cols
    factor=6.9/sum(widths);widths=[x*factor for x in widths]
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for j,vals in enumerate(data):
        row=t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if j==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for k,(cell,txt) in enumerate(zip(row.cells,vals)):
            cell.width=Inches(widths[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if j==0 else 'FFFFFF');pr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
            p.paragraph_format.line_spacing=1.04
            if cols>=3:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,txt)
            for rr in p.runs:rr.font.size=Pt(10);rr.bold=(j==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=md.splitlines();i=0;eq=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line==r'\[':
        while i<len(lines) and lines[i].strip()!=r'\]':i+=1
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p._p.append(copy.deepcopy(equations[eq]));eq+=1
        i+=1;continue
    if line.startswith('#'):
        level=len(line)-len(line.lstrip('#'));s=line.lstrip('#').strip()
        s=re.sub(r'^\d+\.\s*','',s);s=re.sub(r'[^\w\s]',' ',s);s=re.sub(r'\s+',' ',s)
        if level==1:s='Kobon triangle research review'
        doc.add_paragraph(s,style='Title' if level==1 else f'Heading {min(level-1,3)}')
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i]);i+=1
        table(rows);continue
    if line.startswith('* ') or re.match(r'^\d+\. ',line):
        ordered=bool(re.match(r'^\d+\. ',line));txt=re.sub(r'^(\* |\d+\. )','',line)
        p=doc.add_paragraph(style='List Number' if ordered else 'List Bullet');inline(p,txt)
        i+=1;continue
    block=[]
    while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','\\[')):
        block.append(lines[i].lstrip('> ').strip());i+=1
    p=doc.add_paragraph();inline(p,' '.join(block))
    if ' '.join(block).endswith(':'):p.paragraph_format.keep_with_next=True
doc.core_properties.title='Kobon triangle research review'
doc.core_properties.subject='Verified constructions evidence attribution and remaining proof obligations'
doc.core_properties.author='Alejandro Zarzuelo Urdiales'
for root in (doc.styles.element,doc._element):
    for border in list(root.xpath('.//w:pBdr')):border.getparent().remove(border)
dest=ROOT/'research/six-hour-2026-09-21/Kobon_Research_Review.docx'
doc.save(dest);print(dest)
